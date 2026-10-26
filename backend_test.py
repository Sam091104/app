#!/usr/bin/env python3
"""
Session-5d and Session-5e Backend Testing
Tests season management API (CRUD, archive/activate, overview, Excel export) and only_singles_count toggle
"""

import requests
import json
import os
import uuid
from pymongo import MongoClient
from io import BytesIO

# Configuration
BASE_URL = os.getenv('NEXT_PUBLIC_BASE_URL', 'https://pool-stats-3.preview.emergentagent.com')
API_URL = f"{BASE_URL}/api"
ADMIN_HEADER = {'x-admin-code': 'ADMIN-9999'}
MONGO_URL = os.getenv('MONGO_URL', 'mongodb://localhost:27017')
DB_NAME = os.getenv('DB_NAME', 'pool_competitie')

print(f"Testing against: {API_URL}")
print(f"MongoDB: {MONGO_URL}/{DB_NAME}")
print("=" * 80)

# MongoDB connection
mongo_client = MongoClient(MONGO_URL)
db = mongo_client[DB_NAME]

# Global variables for test data
temp_season_id = None
temp_match_id = None
active_season_id = None
archived_season_id = None
den_telstar_id = None
arena_id = None

def test_step_1_get_seasons():
    """Step 1: GET /api/seasons - verify 2 seasons (1 active, 1 archived)"""
    print("\n" + "=" * 80)
    print("STEP 1: GET /api/seasons - List all seasons")
    print("=" * 80)
    
    try:
        resp = requests.get(f"{API_URL}/seasons")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        seasons = resp.json()
        print(f"✅ Got {len(seasons)} seasons")
        
        # Verify we have exactly 2 seasons
        if len(seasons) != 2:
            print(f"❌ FAILED: Expected 2 seasons, got {len(seasons)}")
            return False
        
        # Find active and archived seasons
        active_seasons = [s for s in seasons if s.get('status') == 'active' or s.get('active') == True]
        archived_seasons = [s for s in seasons if s.get('status') == 'archived']
        
        print(f"Active seasons: {len(active_seasons)}")
        print(f"Archived seasons: {len(archived_seasons)}")
        
        if len(active_seasons) != 1:
            print(f"❌ FAILED: Expected 1 active season, got {len(active_seasons)}")
            return False
        
        if len(archived_seasons) != 1:
            print(f"❌ FAILED: Expected 1 archived season, got {len(archived_seasons)}")
            return False
        
        # Verify active season is 2026-2027 with 68 matches
        active = active_seasons[0]
        global active_season_id
        active_season_id = active['id']
        
        print(f"\nActive season: {active['name']}")
        print(f"  ID: {active['id']}")
        print(f"  Status: {active.get('status', 'N/A')}")
        print(f"  Match count: {active.get('match_count', 0)}")
        print(f"  Played count: {active.get('played_count', 0)}")
        print(f"  Has snapshot: {active.get('has_snapshot', False)}")
        print(f"  Only singles count: {active.get('only_singles_count', False)}")
        
        if active['name'] != '2026-2027':
            print(f"❌ FAILED: Expected active season '2026-2027', got '{active['name']}'")
            return False
        
        if active.get('match_count', 0) != 68:
            print(f"❌ FAILED: Expected 68 matches, got {active.get('match_count', 0)}")
            return False
        
        # Verify archived season is 2024-2025
        archived = archived_seasons[0]
        global archived_season_id
        archived_season_id = archived['id']
        
        print(f"\nArchived season: {archived['name']}")
        print(f"  ID: {archived['id']}")
        print(f"  Status: {archived.get('status', 'N/A')}")
        print(f"  Has snapshot: {archived.get('has_snapshot', False)}")
        
        if archived['name'] != '2024-2025':
            print(f"❌ FAILED: Expected archived season '2024-2025', got '{archived['name']}'")
            return False
        
        # Verify all required fields are present
        required_fields = ['id', 'name', 'status', 'active', 'match_count', 'played_count', 'has_snapshot', 'only_singles_count']
        for season in seasons:
            for field in required_fields:
                if field not in season:
                    print(f"❌ FAILED: Missing field '{field}' in season {season.get('name', 'unknown')}")
                    return False
        
        print("\n✅ STEP 1 PASSED: GET /api/seasons returns correct data")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_2_get_season_overview():
    """Step 2: GET /api/seasons/:id/overview - test snapshot vs live"""
    print("\n" + "=" * 80)
    print("STEP 2: GET /api/seasons/:id/overview - Season overview (snapshot vs live)")
    print("=" * 80)
    
    try:
        # Test archived season (should return snapshot)
        print(f"\nTesting archived season overview (2024-2025)...")
        resp = requests.get(f"{API_URL}/seasons/{archived_season_id}/overview")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        data = resp.json()
        print(f"Source: {data.get('source', 'N/A')}")
        print(f"Standings count: {len(data.get('standings', []))}")
        print(f"Player stats count: {len(data.get('player_stats', []))}")
        print(f"Matches count: {len(data.get('matches', []))}")
        
        if data.get('source') != 'snapshot':
            print(f"❌ FAILED: Expected source='snapshot', got '{data.get('source')}'")
            return False
        
        if len(data.get('standings', [])) != 5:
            print(f"❌ FAILED: Expected 5 standings rows, got {len(data.get('standings', []))}")
            return False
        
        if len(data.get('player_stats', [])) != 27:
            print(f"❌ FAILED: Expected 27 player stats, got {len(data.get('player_stats', []))}")
            return False
        
        if len(data.get('matches', [])) != 0:
            print(f"❌ FAILED: Expected 0 matches (snapshot), got {len(data.get('matches', []))}")
            return False
        
        print("✅ Archived season returns snapshot correctly")
        
        # Test active season (should return live)
        print(f"\nTesting active season overview (2026-2027)...")
        resp = requests.get(f"{API_URL}/seasons/{active_season_id}/overview")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        data = resp.json()
        print(f"Source: {data.get('source', 'N/A')}")
        print(f"Standings count: {len(data.get('standings', []))}")
        print(f"Matches count: {len(data.get('matches', []))}")
        
        if data.get('source') != 'live':
            print(f"❌ FAILED: Expected source='live', got '{data.get('source')}'")
            return False
        
        if len(data.get('standings', [])) != 6:
            print(f"❌ FAILED: Expected 6 standings rows, got {len(data.get('standings', []))}")
            return False
        
        if len(data.get('matches', [])) != 68:
            print(f"❌ FAILED: Expected 68 matches, got {len(data.get('matches', []))}")
            return False
        
        print("✅ Active season returns live data correctly")
        
        # Test unknown season ID (should return 404)
        print(f"\nTesting unknown season ID...")
        unknown_id = str(uuid.uuid4())
        resp = requests.get(f"{API_URL}/seasons/{unknown_id}/overview")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 404:
            print(f"❌ FAILED: Expected 404, got {resp.status_code}")
            return False
        
        print("✅ Unknown season ID returns 404")
        
        print("\n✅ STEP 2 PASSED: Season overview works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_3_create_season():
    """Step 3: POST /api/admin/seasons - create temp season with validations"""
    print("\n" + "=" * 80)
    print("STEP 3: POST /api/admin/seasons - Create temp season")
    print("=" * 80)
    
    try:
        # Test without admin header (should return 401)
        print("\nTesting without admin header...")
        resp = requests.post(f"{API_URL}/admin/seasons", json={'name': 'Test'})
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 401:
            print(f"❌ FAILED: Expected 401, got {resp.status_code}")
            return False
        
        print("✅ Without admin header returns 401")
        
        # Test empty body (should return 400)
        print("\nTesting empty body...")
        resp = requests.post(f"{API_URL}/admin/seasons", json={}, headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 400:
            print(f"❌ FAILED: Expected 400, got {resp.status_code}")
            return False
        
        print("✅ Empty body returns 400")
        
        # Create temp season
        print("\nCreating temp season 'Test Seizoen Zeta'...")
        season_data = {
            'name': 'Test Seizoen Zeta',
            'start_date': '2027-09-01',
            'end_date': '2028-06-30',
            'description': 'Temporary test season'
        }
        resp = requests.post(f"{API_URL}/admin/seasons", json=season_data, headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        data = resp.json()
        global temp_season_id
        temp_season_id = data['season']['id']
        
        print(f"✅ Created temp season: {data['season']['name']}")
        print(f"  ID: {temp_season_id}")
        print(f"  Status: {data['season'].get('status', 'N/A')}")
        print(f"  Active: {data['season'].get('active', False)}")
        
        if data['season'].get('status') != 'planned':
            print(f"❌ FAILED: Expected status='planned', got '{data['season'].get('status')}'")
            return False
        
        if data['season'].get('active') != False:
            print(f"❌ FAILED: Expected active=False, got {data['season'].get('active')}")
            return False
        
        # Test duplicate name (should return 409)
        print("\nTesting duplicate name...")
        resp = requests.post(f"{API_URL}/admin/seasons", json=season_data, headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 409:
            print(f"❌ FAILED: Expected 409, got {resp.status_code}")
            return False
        
        print("✅ Duplicate name returns 409")
        
        print("\n✅ STEP 3 PASSED: Season creation works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_4_update_season():
    """Step 4: PATCH /api/admin/seasons/:id - update temp season with validations"""
    print("\n" + "=" * 80)
    print("STEP 4: PATCH /api/admin/seasons/:id - Update temp season")
    print("=" * 80)
    
    try:
        # Test unknown ID (should return 404)
        print("\nTesting unknown season ID...")
        unknown_id = str(uuid.uuid4())
        resp = requests.patch(f"{API_URL}/admin/seasons/{unknown_id}", json={'name': 'Test'}, headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 404:
            print(f"❌ FAILED: Expected 404, got {resp.status_code}")
            return False
        
        print("✅ Unknown season ID returns 404")
        
        # Test empty body (should return 400)
        print("\nTesting empty body...")
        resp = requests.patch(f"{API_URL}/admin/seasons/{temp_season_id}", json={}, headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 400:
            print(f"❌ FAILED: Expected 400, got {resp.status_code}")
            return False
        
        print("✅ Empty body returns 400")
        
        # Update name and description
        print("\nUpdating temp season name and description...")
        update_data = {
            'name': 'Test Seizoen Zeta 2',
            'description': 'Updated description'
        }
        resp = requests.patch(f"{API_URL}/admin/seasons/{temp_season_id}", json=update_data, headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        data = resp.json()
        print(f"✅ Updated season: {data['season']['name']}")
        print(f"  Description: {data['season'].get('description', 'N/A')}")
        
        if data['season']['name'] != 'Test Seizoen Zeta 2':
            print(f"❌ FAILED: Expected name='Test Seizoen Zeta 2', got '{data['season']['name']}'")
            return False
        
        # Test duplicate name (should return 409)
        print("\nTesting duplicate name '2026-2027'...")
        resp = requests.patch(f"{API_URL}/admin/seasons/{temp_season_id}", json={'name': '2026-2027'}, headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 409:
            print(f"❌ FAILED: Expected 409, got {resp.status_code}")
            return False
        
        print("✅ Duplicate name returns 409")
        
        print("\n✅ STEP 4 PASSED: Season update works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_5_archive_season():
    """Step 5: POST /api/admin/seasons/:id/archive - archive temp season"""
    print("\n" + "=" * 80)
    print("STEP 5: POST /api/admin/seasons/:id/archive - Archive temp season")
    print("=" * 80)
    
    try:
        # Archive temp season
        print(f"\nArchiving temp season {temp_season_id}...")
        resp = requests.post(f"{API_URL}/admin/seasons/{temp_season_id}/archive", headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        data = resp.json()
        print(f"✅ Archived season")
        print(f"  Champion: {data.get('champion', 'N/A')}")
        print(f"  No active season: {data.get('no_active_season', False)}")
        
        if data.get('champion') is not None:
            print(f"❌ FAILED: Expected champion=null, got '{data.get('champion')}'")
            return False
        
        if data.get('no_active_season') != False:
            print(f"❌ FAILED: Expected no_active_season=false, got {data.get('no_active_season')}")
            return False
        
        # Verify season status changed to archived
        print("\nVerifying season status...")
        resp = requests.get(f"{API_URL}/seasons")
        seasons = resp.json()
        temp_season = next((s for s in seasons if s['id'] == temp_season_id), None)
        
        if not temp_season:
            print(f"❌ FAILED: Temp season not found in list")
            return False
        
        print(f"  Status: {temp_season.get('status', 'N/A')}")
        print(f"  Has snapshot: {temp_season.get('has_snapshot', False)}")
        
        if temp_season.get('status') != 'archived':
            print(f"❌ FAILED: Expected status='archived', got '{temp_season.get('status')}'")
            return False
        
        # Test archiving already archived season (should return 400)
        print("\nTesting archiving already archived season...")
        resp = requests.post(f"{API_URL}/admin/seasons/{temp_season_id}/archive", headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 400:
            print(f"❌ FAILED: Expected 400, got {resp.status_code}")
            return False
        
        print("✅ Archiving already archived season returns 400")
        
        print("\n✅ STEP 5 PASSED: Season archive works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_6_activate_season():
    """Step 6: POST /api/admin/seasons/:id/activate - ONLY test 404 and 401"""
    print("\n" + "=" * 80)
    print("STEP 6: POST /api/admin/seasons/:id/activate - Test validations ONLY")
    print("=" * 80)
    print("⚠️  CRITICAL: NOT activating any real season to preserve data")
    
    try:
        # Test without admin header (should return 401)
        print("\nTesting without admin header...")
        unknown_id = str(uuid.uuid4())
        resp = requests.post(f"{API_URL}/admin/seasons/{unknown_id}/activate")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 401:
            print(f"❌ FAILED: Expected 401, got {resp.status_code}")
            return False
        
        print("✅ Without admin header returns 401")
        
        # Test unknown ID (should return 404)
        print("\nTesting unknown season ID...")
        resp = requests.post(f"{API_URL}/admin/seasons/{unknown_id}/activate", headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 404:
            print(f"❌ FAILED: Expected 404, got {resp.status_code}")
            return False
        
        print("✅ Unknown season ID returns 404")
        
        print("\n✅ STEP 6 PASSED: Activate endpoint validations work correctly")
        print("⚠️  Did NOT activate any real season (safety preserved)")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_7_export_season():
    """Step 7: GET /api/admin/seasons/:id/export - test Excel export"""
    print("\n" + "=" * 80)
    print("STEP 7: GET /api/admin/seasons/:id/export - Excel export")
    print("=" * 80)
    
    try:
        # Export active season
        print(f"\nExporting active season (2026-2027)...")
        resp = requests.get(f"{API_URL}/admin/seasons/{active_season_id}/export", headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text[:200]}")
            return False
        
        content_type = resp.headers.get('content-type', '')
        print(f"Content-Type: {content_type}")
        
        if 'spreadsheetml' not in content_type:
            print(f"❌ FAILED: Expected content-type to contain 'spreadsheetml', got '{content_type}'")
            return False
        
        # Verify it's a valid xlsx file (starts with 'PK')
        body = resp.content
        print(f"Body size: {len(body)} bytes")
        
        if not body.startswith(b'PK'):
            print(f"❌ FAILED: Excel file should start with 'PK', got '{body[:10]}'")
            return False
        
        print("✅ Excel file format valid (starts with 'PK')")
        
        # Try to parse with openpyxl if available
        try:
            import openpyxl
            wb = openpyxl.load_workbook(BytesIO(body))
            sheet_names = wb.sheetnames
            print(f"Sheet names: {sheet_names}")
            
            expected_sheets = ['Info', 'Ploegstand', 'Individueel klassement', 'Wedstrijden', 'Partijen']
            for expected in expected_sheets:
                if expected not in sheet_names:
                    print(f"❌ FAILED: Missing sheet '{expected}'")
                    return False
            
            print("✅ All expected sheets present")
            
            # Verify Wedstrijden sheet has 68 data rows
            ws = wb['Wedstrijden']
            data_rows = ws.max_row - 1  # Subtract header row
            print(f"Wedstrijden sheet: {data_rows} data rows")
            
            if data_rows != 68:
                print(f"❌ FAILED: Expected 68 data rows in Wedstrijden, got {data_rows}")
                return False
            
            print("✅ Wedstrijden sheet has 68 data rows")
            
        except ImportError:
            print("⚠️  openpyxl not available, installing...")
            import subprocess
            subprocess.run(['pip', 'install', 'openpyxl'], check=True)
            
            # Retry
            import openpyxl
            wb = openpyxl.load_workbook(BytesIO(body))
            sheet_names = wb.sheetnames
            print(f"Sheet names: {sheet_names}")
            
            expected_sheets = ['Info', 'Ploegstand', 'Individueel klassement', 'Wedstrijden', 'Partijen']
            for expected in expected_sheets:
                if expected not in sheet_names:
                    print(f"❌ FAILED: Missing sheet '{expected}'")
                    return False
            
            print("✅ All expected sheets present")
            
            # Verify Wedstrijden sheet has 68 data rows
            ws = wb['Wedstrijden']
            data_rows = ws.max_row - 1  # Subtract header row
            print(f"Wedstrijden sheet: {data_rows} data rows")
            
            if data_rows != 68:
                print(f"❌ FAILED: Expected 68 data rows in Wedstrijden, got {data_rows}")
                return False
            
            print("✅ Wedstrijden sheet has 68 data rows")
        
        print("\n✅ STEP 7 PASSED: Excel export works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_8_delete_season():
    """Step 8: DELETE /api/admin/seasons/:id - test delete rules"""
    print("\n" + "=" * 80)
    print("STEP 8: DELETE /api/admin/seasons/:id - Delete season")
    print("=" * 80)
    
    try:
        # Test deleting active season (should return 400)
        print(f"\nTesting delete active season (should fail)...")
        resp = requests.delete(f"{API_URL}/admin/seasons/{active_season_id}", headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 400:
            print(f"❌ FAILED: Expected 400, got {resp.status_code}")
            return False
        
        print("✅ Deleting active season returns 400")
        
        # Test unknown ID (should return 404)
        print("\nTesting unknown season ID...")
        unknown_id = str(uuid.uuid4())
        resp = requests.delete(f"{API_URL}/admin/seasons/{unknown_id}", headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 404:
            print(f"❌ FAILED: Expected 404, got {resp.status_code}")
            return False
        
        print("✅ Unknown season ID returns 404")
        
        # Note: We'll delete the temp season in step 15 after testing only_singles_count
        
        print("\n✅ STEP 8 PASSED: Delete validations work correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_9_season_filtering():
    """Step 9: Test season filtering on matches and standings"""
    print("\n" + "=" * 80)
    print("STEP 9: Season filtering - GET /api/matches and /api/standings")
    print("=" * 80)
    
    try:
        # Test GET /api/matches?season_id=all
        print("\nTesting GET /api/matches?season_id=all...")
        resp = requests.get(f"{API_URL}/matches?season_id=all")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            return False
        
        matches = resp.json()
        print(f"✅ Got {len(matches)} matches")
        
        if len(matches) != 68:
            print(f"❌ FAILED: Expected 68 matches, got {len(matches)}")
            return False
        
        # Test GET /api/matches?season_id=<random uuid> (should return empty)
        print("\nTesting GET /api/matches?season_id=<random uuid>...")
        random_id = str(uuid.uuid4())
        resp = requests.get(f"{API_URL}/matches?season_id={random_id}")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            return False
        
        matches = resp.json()
        print(f"✅ Got {len(matches)} matches")
        
        if len(matches) != 0:
            print(f"❌ FAILED: Expected 0 matches, got {len(matches)}")
            return False
        
        # Test GET /api/standings?season_id=<archived>
        print(f"\nTesting GET /api/standings?season_id={archived_season_id}...")
        resp = requests.get(f"{API_URL}/standings?season_id={archived_season_id}")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            return False
        
        standings = resp.json()
        print(f"✅ Got {len(standings)} standings (computed live for archived season)")
        
        print("\n✅ STEP 9 PASSED: Season filtering works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_10_insert_test_match():
    """Step 10: Insert test match directly in Mongo for temp season"""
    print("\n" + "=" * 80)
    print("STEP 10: Insert test match in Mongo for temp season")
    print("=" * 80)
    
    try:
        # Get team IDs
        print("\nGetting team IDs...")
        resp = requests.get(f"{API_URL}/teams")
        teams = resp.json()
        
        global den_telstar_id, arena_id
        den_telstar = next((t for t in teams if t['name'] == 'Den Telstar'), None)
        arena = next((t for t in teams if t['name'] == 'Arena'), None)
        
        if not den_telstar or not arena:
            print(f"❌ FAILED: Could not find Den Telstar or Arena teams")
            return False
        
        den_telstar_id = den_telstar['id']
        arena_id = arena['id']
        
        print(f"Den Telstar ID: {den_telstar_id}")
        print(f"Arena ID: {arena_id}")
        
        # Create match document
        global temp_match_id
        temp_match_id = str(uuid.uuid4())
        
        # 8 enkel games (all 2-0 home wins) + 2 dubbel games (all 0-2 away wins)
        games = []
        for i in range(8):
            games.append({
                'type': 'enkel',
                'home_players': [],
                'away_players': [],
                'home_frames': 2,
                'away_frames': 0,
                'forfait_home': False,
                'forfait_away': False
            })
        for i in range(2):
            games.append({
                'type': 'dubbel',
                'home_players': [],
                'away_players': [],
                'home_frames': 0,
                'away_frames': 2,
                'forfait_home': False,
                'forfait_away': False
            })
        
        match_doc = {
            'id': temp_match_id,
            'season_id': temp_season_id,
            'week': 1,
            'date': '2027-10-01',
            'home_team_id': den_telstar_id,
            'away_team_id': arena_id,
            'status': 'confirmed',
            'points_home': 0,
            'points_away': 0,
            'games': games,
            'created_at': '2027-09-01T00:00:00.000Z'
        }
        
        print(f"\nInserting match {temp_match_id}...")
        print(f"  Season: {temp_season_id}")
        print(f"  Home: Den Telstar ({den_telstar_id})")
        print(f"  Away: Arena ({arena_id})")
        print(f"  Games: 8 enkel (2-0 home) + 2 dubbel (0-2 away)")
        
        result = db.matches.insert_one(match_doc)
        print(f"✅ Inserted match: {result.inserted_id}")
        
        # Verify match was inserted
        match = db.matches.find_one({'id': temp_match_id})
        if not match:
            print(f"❌ FAILED: Match not found in database")
            return False
        
        print(f"✅ Match verified in database")
        
        print("\n✅ STEP 10 PASSED: Test match inserted successfully")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_11_only_singles_count_true():
    """Step 11: PATCH only_singles_count:true - verify recalculation"""
    print("\n" + "=" * 80)
    print("STEP 11: PATCH only_singles_count:true - Recalculation test")
    print("=" * 80)
    
    try:
        # Enable only_singles_count
        print(f"\nEnabling only_singles_count for temp season...")
        resp = requests.patch(
            f"{API_URL}/admin/seasons/{temp_season_id}",
            json={'only_singles_count': True},
            headers=ADMIN_HEADER
        )
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        data = resp.json()
        print(f"✅ Updated season")
        print(f"  only_singles_count: {data['season'].get('only_singles_count', False)}")
        
        if 'recalculated' in data:
            print(f"  Recalculated: checked={data['recalculated'].get('checked', 0)}, updated={data['recalculated'].get('updated', 0)}")
            
            if data['recalculated'].get('checked', 0) != 1:
                print(f"❌ FAILED: Expected recalculated.checked=1, got {data['recalculated'].get('checked', 0)}")
                return False
            
            if data['recalculated'].get('updated', 0) != 1:
                print(f"❌ FAILED: Expected recalculated.updated=1, got {data['recalculated'].get('updated', 0)}")
                return False
        
        # Verify match points were recalculated (only enkels count)
        # 8 enkel games (2-0 home) = 8*2 = 16 points home, 0 points away
        # 2 dubbel games (0-2 away) = excluded
        print(f"\nVerifying match points (only enkels)...")
        match = db.matches.find_one({'id': temp_match_id})
        
        print(f"  points_home: {match.get('points_home', 0)}")
        print(f"  points_away: {match.get('points_away', 0)}")
        
        if match.get('points_home', 0) != 16:
            print(f"❌ FAILED: Expected points_home=16, got {match.get('points_home', 0)}")
            return False
        
        if match.get('points_away', 0) != 0:
            print(f"❌ FAILED: Expected points_away=0, got {match.get('points_away', 0)}")
            return False
        
        print("✅ Match points correct (16-0, dubbels excluded)")
        
        # Verify standings
        print(f"\nVerifying standings...")
        resp = requests.get(f"{API_URL}/standings?season_id={temp_season_id}")
        standings = resp.json()
        
        den_telstar_standing = next((s for s in standings if s['team_name'] == 'Den Telstar'), None)
        arena_standing = next((s for s in standings if s['team_name'] == 'Arena'), None)
        
        if den_telstar_standing:
            print(f"Den Telstar:")
            print(f"  points_for: {den_telstar_standing.get('points_for', 0)}")
            print(f"  points_against: {den_telstar_standing.get('points_against', 0)}")
            print(f"  frames_won: {den_telstar_standing.get('frames_won', 0)}")
            print(f"  frames_lost: {den_telstar_standing.get('frames_lost', 0)}")
            print(f"  won: {den_telstar_standing.get('won', 0)}")
            
            if den_telstar_standing.get('points_for', 0) != 16:
                print(f"❌ FAILED: Expected Den Telstar points_for=16, got {den_telstar_standing.get('points_for', 0)}")
                return False
            
            if den_telstar_standing.get('points_against', 0) != 0:
                print(f"❌ FAILED: Expected Den Telstar points_against=0, got {den_telstar_standing.get('points_against', 0)}")
                return False
            
            if den_telstar_standing.get('frames_won', 0) != 16:
                print(f"❌ FAILED: Expected Den Telstar frames_won=16, got {den_telstar_standing.get('frames_won', 0)}")
                return False
            
            if den_telstar_standing.get('frames_lost', 0) != 0:
                print(f"❌ FAILED: Expected Den Telstar frames_lost=0, got {den_telstar_standing.get('frames_lost', 0)}")
                return False
            
            if den_telstar_standing.get('won', 0) != 1:
                print(f"❌ FAILED: Expected Den Telstar won=1, got {den_telstar_standing.get('won', 0)}")
                return False
        
        if arena_standing:
            print(f"Arena:")
            print(f"  points_for: {arena_standing.get('points_for', 0)}")
            
            if arena_standing.get('points_for', 0) != 0:
                print(f"❌ FAILED: Expected Arena points_for=0, got {arena_standing.get('points_for', 0)}")
                return False
        
        print("✅ Standings correct (only enkels counted)")
        
        print("\n✅ STEP 11 PASSED: only_singles_count=true works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_12_only_singles_count_false():
    """Step 12: PATCH only_singles_count:false - verify recalculation"""
    print("\n" + "=" * 80)
    print("STEP 12: PATCH only_singles_count:false - Recalculation test")
    print("=" * 80)
    
    try:
        # Disable only_singles_count
        print(f"\nDisabling only_singles_count for temp season...")
        resp = requests.patch(
            f"{API_URL}/admin/seasons/{temp_season_id}",
            json={'only_singles_count': False},
            headers=ADMIN_HEADER
        )
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        data = resp.json()
        print(f"✅ Updated season")
        print(f"  only_singles_count: {data['season'].get('only_singles_count', False)}")
        
        if 'recalculated' in data:
            print(f"  Recalculated: checked={data['recalculated'].get('checked', 0)}, updated={data['recalculated'].get('updated', 0)}")
            
            if data['recalculated'].get('updated', 0) != 1:
                print(f"❌ FAILED: Expected recalculated.updated=1, got {data['recalculated'].get('updated', 0)}")
                return False
        
        # Verify match points were recalculated (all games count)
        # 8 enkel games (2-0 home) = 8*2 = 16 points home, 0 points away
        # 2 dubbel games (0-2 away) = 0 points home, 2*2 = 4 points away
        # Total: 16 home, 4 away
        print(f"\nVerifying match points (all games)...")
        match = db.matches.find_one({'id': temp_match_id})
        
        print(f"  points_home: {match.get('points_home', 0)}")
        print(f"  points_away: {match.get('points_away', 0)}")
        
        if match.get('points_home', 0) != 16:
            print(f"❌ FAILED: Expected points_home=16, got {match.get('points_home', 0)}")
            return False
        
        if match.get('points_away', 0) != 4:
            print(f"❌ FAILED: Expected points_away=4, got {match.get('points_away', 0)}")
            return False
        
        print("✅ Match points correct (16-4, dubbels included)")
        
        # Verify standings
        print(f"\nVerifying standings...")
        resp = requests.get(f"{API_URL}/standings?season_id={temp_season_id}")
        standings = resp.json()
        
        den_telstar_standing = next((s for s in standings if s['team_name'] == 'Den Telstar'), None)
        
        if den_telstar_standing:
            print(f"Den Telstar:")
            print(f"  points_for: {den_telstar_standing.get('points_for', 0)}")
            print(f"  points_against: {den_telstar_standing.get('points_against', 0)}")
            print(f"  frames_lost: {den_telstar_standing.get('frames_lost', 0)}")
            
            if den_telstar_standing.get('points_for', 0) != 16:
                print(f"❌ FAILED: Expected Den Telstar points_for=16, got {den_telstar_standing.get('points_for', 0)}")
                return False
            
            if den_telstar_standing.get('points_against', 0) != 4:
                print(f"❌ FAILED: Expected Den Telstar points_against=4, got {den_telstar_standing.get('points_against', 0)}")
                return False
            
            if den_telstar_standing.get('frames_lost', 0) != 4:
                print(f"❌ FAILED: Expected Den Telstar frames_lost=4, got {den_telstar_standing.get('frames_lost', 0)}")
                return False
        
        print("✅ Standings correct (all games counted)")
        
        print("\n✅ STEP 12 PASSED: only_singles_count=false works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_13_only_singles_archived():
    """Step 13: PATCH only_singles_count on archived season - verify 400"""
    print("\n" + "=" * 80)
    print("STEP 13: PATCH only_singles_count on archived season - Should fail")
    print("=" * 80)
    
    try:
        # Try to change only_singles_count on archived season
        print(f"\nTrying to change only_singles_count on archived season 2024-2025...")
        resp = requests.patch(
            f"{API_URL}/admin/seasons/{archived_season_id}",
            json={'only_singles_count': True},
            headers=ADMIN_HEADER
        )
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 400:
            print(f"❌ FAILED: Expected 400, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        print("✅ Changing only_singles_count on archived season returns 400")
        
        # Test PATCH with same value as current (should return 400 if nothing else changes)
        print(f"\nTrying to PATCH temp season with same only_singles_count value...")
        resp = requests.patch(
            f"{API_URL}/admin/seasons/{temp_season_id}",
            json={'only_singles_count': False},  # Already false from step 12
            headers=ADMIN_HEADER
        )
        print(f"Status: {resp.status_code}")
        
        # This might return 200 or 400 depending on implementation
        # If it returns 400, it should have a message about no changes
        if resp.status_code == 400:
            print("✅ PATCH with same value returns 400 (no changes)")
        elif resp.status_code == 200:
            data = resp.json()
            if 'recalculated' in data and data['recalculated'].get('updated', 0) == 0:
                print("✅ PATCH with same value returns 200 but no recalculation")
            else:
                print("⚠️  PATCH with same value returns 200 (implementation allows this)")
        
        print("\n✅ STEP 13 PASSED: Archived season protection works correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_14_players_season_filter():
    """Step 14: GET /api/players?season_id=<zeta> - verify no crash"""
    print("\n" + "=" * 80)
    print("STEP 14: GET /api/players?season_id=<zeta> - No crash test")
    print("=" * 80)
    
    try:
        # Get players for temp season (should not crash even with empty player IDs in games)
        print(f"\nGetting players for temp season...")
        resp = requests.get(f"{API_URL}/players?season_id={temp_season_id}")
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        players = resp.json()
        print(f"✅ Got {len(players)} players (no crash)")
        
        print("\n✅ STEP 14 PASSED: Players endpoint handles empty player IDs correctly")
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_step_15_cleanup():
    """Step 15: Cleanup - delete temp season and verify final state"""
    print("\n" + "=" * 80)
    print("STEP 15: Cleanup - Delete temp season and verify final state")
    print("=" * 80)
    
    try:
        # Delete temp season
        print(f"\nDeleting temp season {temp_season_id}...")
        resp = requests.delete(f"{API_URL}/admin/seasons/{temp_season_id}", headers=ADMIN_HEADER)
        print(f"Status: {resp.status_code}")
        
        if resp.status_code != 200:
            print(f"❌ FAILED: Expected 200, got {resp.status_code}")
            print(f"Response: {resp.text}")
            return False
        
        data = resp.json()
        print(f"✅ Deleted temp season")
        print(f"  matches_deleted: {data.get('matches_deleted', 0)}")
        
        if data.get('matches_deleted', 0) != 1:
            print(f"❌ FAILED: Expected matches_deleted=1, got {data.get('matches_deleted', 0)}")
            return False
        
        # Verify temp season is gone
        print(f"\nVerifying temp season is deleted...")
        resp = requests.get(f"{API_URL}/seasons")
        seasons = resp.json()
        
        temp_season = next((s for s in seasons if s['id'] == temp_season_id), None)
        if temp_season:
            print(f"❌ FAILED: Temp season still exists in list")
            return False
        
        print("✅ Temp season deleted from list")
        
        # Verify only 2 seasons remain
        print(f"\nVerifying final state...")
        print(f"Total seasons: {len(seasons)}")
        
        if len(seasons) != 2:
            print(f"❌ FAILED: Expected 2 seasons, got {len(seasons)}")
            return False
        
        # Verify active season is still 2026-2027 with 68 matches
        active_seasons = [s for s in seasons if s.get('status') == 'active' or s.get('active') == True]
        if len(active_seasons) != 1:
            print(f"❌ FAILED: Expected 1 active season, got {len(active_seasons)}")
            return False
        
        active = active_seasons[0]
        print(f"\nActive season: {active['name']}")
        print(f"  Match count: {active.get('match_count', 0)}")
        
        if active['name'] != '2026-2027':
            print(f"❌ FAILED: Expected active season '2026-2027', got '{active['name']}'")
            return False
        
        if active.get('match_count', 0) != 68:
            print(f"❌ FAILED: Expected 68 matches, got {active.get('match_count', 0)}")
            return False
        
        # Verify archived season is still 2024-2025
        archived_seasons = [s for s in seasons if s.get('status') == 'archived']
        if len(archived_seasons) != 1:
            print(f"❌ FAILED: Expected 1 archived season, got {len(archived_seasons)}")
            return False
        
        archived = archived_seasons[0]
        print(f"\nArchived season: {archived['name']}")
        
        if archived['name'] != '2024-2025':
            print(f"❌ FAILED: Expected archived season '2024-2025', got '{archived['name']}'")
            return False
        
        # Verify GET /api/matches?season_id=all returns 68
        print(f"\nVerifying total matches...")
        resp = requests.get(f"{API_URL}/matches?season_id=all")
        matches = resp.json()
        print(f"Total matches: {len(matches)}")
        
        if len(matches) != 68:
            print(f"❌ FAILED: Expected 68 matches, got {len(matches)}")
            return False
        
        print("\n✅ STEP 15 PASSED: Cleanup successful, final state verified")
        print("\n" + "=" * 80)
        print("FINAL STATE VERIFICATION:")
        print("  ✅ 2 seasons remain (2026-2027 active, 2024-2025 archived)")
        print("  ✅ Active season has 68 matches")
        print("  ✅ Temp season and its match deleted")
        print("  ✅ No data corruption")
        print("=" * 80)
        return True
        
    except Exception as e:
        print(f"❌ FAILED: Exception: {e}")
        import traceback
        traceback.print_exc()
        return False


def main():
    """Run all tests"""
    print("\n" + "=" * 80)
    print("SESSION-5D AND SESSION-5E BACKEND TESTING")
    print("Season Management API + only_singles_count Toggle")
    print("=" * 80)
    
    results = []
    
    # Run all tests in sequence
    tests = [
        ("Step 1: GET /api/seasons", test_step_1_get_seasons),
        ("Step 2: GET /api/seasons/:id/overview", test_step_2_get_season_overview),
        ("Step 3: POST /api/admin/seasons", test_step_3_create_season),
        ("Step 4: PATCH /api/admin/seasons/:id", test_step_4_update_season),
        ("Step 5: POST /api/admin/seasons/:id/archive", test_step_5_archive_season),
        ("Step 6: POST /api/admin/seasons/:id/activate", test_step_6_activate_season),
        ("Step 7: GET /api/admin/seasons/:id/export", test_step_7_export_season),
        ("Step 8: DELETE /api/admin/seasons/:id", test_step_8_delete_season),
        ("Step 9: Season filtering", test_step_9_season_filtering),
        ("Step 10: Insert test match", test_step_10_insert_test_match),
        ("Step 11: only_singles_count=true", test_step_11_only_singles_count_true),
        ("Step 12: only_singles_count=false", test_step_12_only_singles_count_false),
        ("Step 13: only_singles_count on archived", test_step_13_only_singles_archived),
        ("Step 14: GET /api/players?season_id", test_step_14_players_season_filter),
        ("Step 15: Cleanup", test_step_15_cleanup),
    ]
    
    for test_name, test_func in tests:
        try:
            result = test_func()
            results.append((test_name, result))
        except Exception as e:
            print(f"\n❌ {test_name} FAILED with exception: {e}")
            import traceback
            traceback.print_exc()
            results.append((test_name, False))
    
    # Print summary
    print("\n" + "=" * 80)
    print("TEST SUMMARY")
    print("=" * 80)
    
    passed = sum(1 for _, result in results if result)
    total = len(results)
    
    for test_name, result in results:
        status = "✅ PASSED" if result else "❌ FAILED"
        print(f"{status}: {test_name}")
    
    print("=" * 80)
    print(f"TOTAL: {passed}/{total} tests passed")
    print("=" * 80)
    
    if passed == total:
        print("\n🎉 ALL TESTS PASSED! 🎉")
        return 0
    else:
        print(f"\n⚠️  {total - passed} test(s) failed")
        return 1


if __name__ == '__main__':
    exit(main())
